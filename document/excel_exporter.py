"""XLSX exporter (openpyxl): typed cells, safe text, validated formulas, tables."""
from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config.constants import DEFAULT_ACCENT_HEX, EXCEL_MAX_COL_WIDTH, EXCEL_TABLE_STYLE
from core.exceptions import FormulaPolicyError
from core.models import Cell, Sheet, WorkbookContentModel
from document.base_exporter import FileExporter
from document.formula_safety import validate_formula, write_text_cell
from document.template_registry import Template

_BAD_TITLE_RE = re.compile(r"[\[\]:*?/\\]")
_HEX_RE = re.compile(r"^[0-9A-Fa-f]{6}$")


def _hex(value: Any, default: str) -> str:
    text = str(value or "").lstrip("#")
    return text.upper() if _HEX_RE.match(text) else default


def _parse_date(text: str) -> date | datetime:
    text = text.strip()
    if len(text) == 10:
        return date.fromisoformat(text)
    return datetime.fromisoformat(text).replace(tzinfo=None)  # Excel has no time zones


class ExcelExporter(FileExporter):
    extension = ".xlsx"
    content_type = WorkbookContentModel
    features = frozenset({"sheets", "typed_cells", "formulas", "tables", "freeze_panes",
                          "number_formats", "rtl"})

    def __init__(self, template: Template) -> None:
        cfg = template.excel
        self._fill = _hex(cfg.get("header_fill"), DEFAULT_ACCENT_HEX)
        self._font_color = _hex(cfg.get("header_font_color"), "FFFFFF")
        self._col_width = float(cfg.get("default_col_width", 18))
        self._freeze = bool(cfg.get("freeze_panes", True))

    # --- FileExporter hooks ---
    def _write(self, content: WorkbookContentModel, path: Path, warnings: list[str]) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        wb.properties.title = content.metadata.title
        wb.properties.creator = content.metadata.author or "Document Builder AI"
        used: set[str] = set()
        sheets = content.sheets or [Sheet("Sheet1", ("",), ())]
        for index, sheet in enumerate(sheets, 1):
            self._sheet(wb, sheet, index, used, content.metadata.language == "ar", warnings)
        wb.save(str(path))

    def _verify(self, path: Path) -> None:
        wb = load_workbook(str(path))
        try:
            if not wb.sheetnames:
                raise ValueError("workbook has no sheets")
        finally:
            wb.close()

    # --- sheets ---
    @staticmethod
    def _unique_title(name: str, used: set[str]) -> str:
        base = (_BAD_TITLE_RE.sub("_", name).strip() or "Sheet")[:31]
        title, n = base, 2
        while title.lower() in used:
            suffix = f" ({n})"
            title, n = base[:31 - len(suffix)] + suffix, n + 1
        used.add(title.lower())
        return title

    def _sheet(self, wb: Workbook, sheet: Sheet, index: int, used: set[str], rtl: bool,
               warnings: list[str]) -> None:
        ws = wb.create_sheet(self._unique_title(sheet.name, used))
        if rtl:
            ws.sheet_view.rightToLeft = True
        ncols = max([len(sheet.columns), *(len(r) for r in sheet.rows)] or [1])
        if ncols > len(sheet.columns):
            warnings.append(f"Sheet '{ws.title}': rows were wider than the header; "
                            "generic column names were added.")
        headers = self._headers(sheet.columns, ncols)
        header_style = (Font(bold=True, color=self._font_color),
                        PatternFill("solid", start_color=self._fill))
        for col, text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            write_text_cell(cell, text)
            cell.font, cell.fill = header_style
        widths = [len(h) for h in headers]
        for r, row in enumerate(sheet.rows, 2):
            for c in range(ncols):
                data = row[c] if c < len(row) else Cell("blank")
                width = self._cell(ws.cell(row=r, column=c + 1), data, warnings,
                                   f"{ws.title}!{get_column_letter(c + 1)}{r}")
                widths[c] = max(widths[c], width)
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = max(
                self._col_width, min(width + 2, EXCEL_MAX_COL_WIDTH))
        if self._freeze and sheet.freeze_header:
            ws.freeze_panes = "A2"
        if sheet.rows:  # Excel tables need at least one data row
            table = Table(displayName=f"Table{index}",
                          ref=f"A1:{get_column_letter(ncols)}{len(sheet.rows) + 1}")
            table.tableStyleInfo = TableStyleInfo(name=EXCEL_TABLE_STYLE, showRowStripes=True)
            ws.add_table(table)

    @staticmethod
    def _headers(columns: tuple[str, ...], ncols: int) -> list[str]:
        """Table headers must be non-empty and unique."""
        out, seen = [], set()
        for i in range(ncols):
            text = (columns[i] if i < len(columns) else "").strip() or f"Column {i + 1}"
            base, n = text, 2
            while text.lower() in seen:
                text, n = f"{base} ({n})", n + 1
            seen.add(text.lower())
            out.append(text)
        return out

    # --- cells ---
    def _cell(self, target: Any, cell: Cell, warnings: list[str], where: str) -> int:
        """Write one typed cell; returns its display width."""
        kind, value = cell.type, cell.value
        try:
            if kind == "blank" or value is None:
                return 0
            if kind == "number":
                if isinstance(value, bool) or not isinstance(value, (int, float)) \
                        or not math.isfinite(value):
                    raise ValueError("not a finite number")
                target.value = value
            elif kind == "boolean":
                if not isinstance(value, bool):
                    raise ValueError("not a boolean")
                target.value = value
            elif kind == "date":
                target.value = _parse_date(str(value))
                if not cell.number_format:
                    target.number_format = ("yyyy-mm-dd" if isinstance(target.value, date)
                                            and not isinstance(target.value, datetime)
                                            else "yyyy-mm-dd hh:mm")
            elif kind == "formula":
                target.value = validate_formula(str(value))
            else:  # "text" and anything unknown: always plain text
                write_text_cell(target, str(value))
        except (ValueError, TypeError, FormulaPolicyError) as exc:
            write_text_cell(target, str(value))  # keep the data, but as inert text
            warnings.append(f"{where}: {kind} value written as text ({exc}).")
            return len(str(value))
        if cell.number_format and kind in ("number", "date"):
            target.number_format = cell.number_format
        return len(str(value))
