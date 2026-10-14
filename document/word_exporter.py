"""DOCX exporter (python-docx) with template-driven styles and RTL support."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from config.constants import DEFAULT_ACCENT_HEX
from core.models import DocumentContentModel, Heading, ListBlock, Paragraph, Table
from document.base_exporter import FileExporter
from document.template_registry import Template
from document.text_utils import clean_text, normalize_table, resolve_direction

# Child order required by the OOXML schema (Word can reject out-of-order elements).
_RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
        "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
        "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
        "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
        "specVanish", "oMath"]
_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
        "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
        "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
        "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap",
        "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId",
        "cnfStyle", "rPr", "sectPr", "pPrChange"]
_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
          "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
          "tblLayout", "tblCellMar", "tblLook"]
_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
         "textDirection", "tcFitText", "vAlign", "hideMark"]


def _put(parent: Any, name: str, order: list[str], **attrs: str) -> Any:
    """Insert (or replace) a child at its schema-correct position."""
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    el = OxmlElement(f"w:{name}")
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), value)
    later = {qn(f"w:{n}") for n in order[order.index(name) + 1:]}
    for i, child in enumerate(parent):
        if child.tag in later:
            parent.insert(i, el)
            return el
    parent.append(el)
    return el


def _hex(value: Any) -> str:
    text = str(value or "").lstrip("#").upper()
    ok = len(text) == 6 and all(c in "0123456789ABCDEF" for c in text)
    return text if ok else DEFAULT_ACCENT_HEX


class WordExporter(FileExporter):
    extension = ".docx"
    content_type = DocumentContentModel
    features = frozenset({"headings", "paragraphs", "lists", "tables", "rtl", "mixed_direction",
                          "heading_numbering", "page_numbers"})

    def __init__(self, template: Template) -> None:
        cfg = template.word
        self._latin = str(cfg.get("font", "Calibri"))
        self._cs = str(cfg.get("arabic_font", "Arial"))
        self._size = float(cfg.get("body_size_pt", 11))
        self._margin = float(cfg.get("margins_cm", 2.5))
        self._numbered = bool(cfg.get("heading_numbering", False))
        self._header_fill = _hex(template.excel.get("header_fill"))

    # --- FileExporter hooks ---
    def _write(self, content: DocumentContentModel, path: Path, warnings: list[str]) -> None:
        doc = Document()
        self._setup(doc, content)
        rtl_doc = content.metadata.language == "ar"
        counters = [0, 0, 0, 0]
        if content.metadata.title:
            self._fill(doc.add_paragraph(style="Title"), content.metadata.title)
        for block in content.blocks:
            if isinstance(block, Heading):
                self._heading(doc, block, counters)
            elif isinstance(block, Paragraph):
                self._fill(doc.add_paragraph(), block.text, block.direction)
            elif isinstance(block, ListBlock):
                self._list(doc, block)
            elif isinstance(block, Table):
                note = self._table(doc, block, rtl_doc)
                if note and note not in warnings:
                    warnings.append(note)
        doc.save(str(path))

    def _verify(self, path: Path) -> None:
        Document(str(path))

    # --- document setup ---
    def _setup(self, doc: Any, content: DocumentContentModel) -> None:
        meta = content.metadata
        doc.core_properties.title = meta.title
        doc.core_properties.author = meta.author or "Document Builder AI"
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4",
                     "Caption", "List Bullet"):
            self._style_fonts(doc.styles[name], self._size if name == "Normal" else None)
        sec = doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(self._margin)
        sec.top_margin = sec.bottom_margin = Cm(self._margin)
        if meta.title:
            self._fill(sec.header.paragraphs[0], meta.title)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._page_field(footer)

    def _style_fonts(self, style: Any, size: float | None) -> None:
        style.font.name = self._latin
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)  # theme attrs override explicit fonts
        fonts.set(qn("w:cs"), self._cs)
        if size:
            style.font.size = Pt(size)
            _put(rpr, "szCs", _RPR, val=str(int(size * 2)))

    @staticmethod
    def _page_field(paragraph: Any) -> None:
        """Insert a PAGE field (complex-field form)."""
        for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
            run = paragraph.add_run()
            if kind:
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), kind)
                run._r.append(fld)
            else:
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = text
                run._r.append(instr)

    # --- blocks ---
    def _fill(self, paragraph: Any, text: str, hint: str = "auto", bold: bool = False) -> bool:
        """Write text into a paragraph; direction is set at paragraph level. Returns is_rtl."""
        text = clean_text(text)
        rtl = resolve_direction(text, hint) == "rtl"
        run = paragraph.add_run(text)
        if bold:
            run.bold = True
        if rtl:
            # bidi flips alignment itself: never force RIGHT/LEFT here.
            _put(paragraph._p.get_or_add_pPr(), "bidi", _PPR)
            run.font.rtl = True
            if bold:
                _put(run._r.get_or_add_rPr(), "bCs", _RPR)
        return rtl

    def _heading(self, doc: Any, block: Heading, counters: list[int]) -> None:
        level = min(max(block.level, 1), 4)
        text = block.text
        if self._numbered:
            counters[level - 1] += 1
            for i in range(level, 4):
                counters[i] = 0
            # No trailing period: it would be displaced in RTL paragraphs.
            text = ".".join(str(n) for n in counters[:level]) + " " + text
        self._fill(doc.add_paragraph(style=f"Heading {level}"), text)

    def _list(self, doc: Any, block: ListBlock) -> None:
        for n, item in enumerate(block.items, 1):
            if not block.ordered:
                self._fill(doc.add_paragraph(style="List Bullet"), item)
                continue
            # Manual numbers: List Number would continue numbering across lists.
            para = doc.add_paragraph()
            rtl = self._fill(para, f"{n}\t{item}" if resolve_direction(item) == "rtl"
                             else f"{n}.\t{item}")
            if not rtl:
                para.paragraph_format.left_indent = Cm(0.9)
                para.paragraph_format.first_line_indent = Cm(-0.9)

    def _table(self, doc: Any, block: Table, rtl_doc: bool) -> str | None:
        headers, rows, note = normalize_table(block.headers, block.rows)
        if block.caption:
            self._fill(doc.add_paragraph(style="Caption"), block.caption)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        if rtl_doc or resolve_direction(" ".join(headers)) == "rtl":
            _put(table._tbl.tblPr, "bidiVisual", _TBLPR)
        head_row = table.rows[0]
        head_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, text in zip(head_row.cells, headers):
            self._fill(cell.paragraphs[0], text, bold=True)
            _put(cell._tc.get_or_add_tcPr(), "shd", _TCPR, val="clear", color="auto",
                 fill=self._header_fill)
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = _white()
        for row in rows:
            for cell, text in zip(table.add_row().cells, row):
                self._fill(cell.paragraphs[0], text)
        doc.add_paragraph()  # spacing after table
        return note


def _white() -> Any:
    from docx.shared import RGBColor
    return RGBColor(0xFF, 0xFF, 0xFF)
