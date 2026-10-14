import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pypdf import PdfReader

from config.constants import ContentKind, OutputFormat
from core.exceptions import (ExportError, FontNotFoundError, FormulaPolicyError,
                             OutputExistsError, UnsupportedContentError)
from core.models import (Cell, DocumentContentModel, Heading, ListBlock, Metadata, Paragraph,
                         Sheet, Table, WorkbookContentModel)
from document.base_exporter import FileExporter
from document.excel_exporter import ExcelExporter
from document.formula_safety import is_formula_like, validate_formula
from document.pdf_exporter_reportlab import PdfExporter
from document.template_registry import Template
from document.word_exporter import WordExporter
from services.document_service import DocumentService, safe_filename, unique_path

TPL = Template("t", "T", "d", ContentKind.DOCUMENT, "document_content.schema.json",
               {"font": "Calibri", "arabic_font": "Arial", "body_size_pt": 11, "margins_cm": 2.5,
                "heading_numbering": True},
               {"header_fill": "1F4E78", "header_font_color": "FFFFFF", "default_col_width": 18,
                "freeze_panes": True},
               {"font": "Helvetica", "margins_cm": 2.0, "header": True, "footer": True})

_FONT_DIRS = [Path("/usr/share/fonts/truetype/dejavu"), Path("C:/Windows/Fonts")]
FONT_DIR = next((d for d in _FONT_DIRS if (d / "DejaVuSans.ttf").is_file()), None)


class FakeShaper:
    """Identity reshape + reversal: exercises wrapping without arabic libs."""
    def reshape(self, text): return text
    def to_visual(self, line, direction): return line[::-1] if direction == "rtl" else line


def doc_content(lang="en"):
    return DocumentContentModel(Metadata("Report", lang, "me"), [
        Heading("Intro", 1), Heading("Sub", 2),
        Paragraph("Hello world"),
        Paragraph("مرحبا بالعالم Mixed text", "auto"),
        ListBlock(("a", "b"), True), ListBlock(("x", "y")),
        Table(("h1", "h2"), (("1", "2"), ("3", "4", "extra")), "Cap"),
    ])


def wb_content(lang="en"):
    rows = (
        (Cell("text", "=1+1"), Cell("number", -5), Cell("formula", "=SUM(B2:B2)")),
        (Cell("text", "  =cmd|x"), Cell("text", "-5"), Cell("formula", '=HYPERLINK("http://x")')),
        (Cell("text", "\t=x"), Cell("text", "@SUM(1)"), Cell("date", "2026-10-14")),
        (Cell("boolean", True), Cell("blank"), Cell("number", 0.25, "0%")),
    )
    return WorkbookContentModel(Metadata("W", lang), [Sheet("Data", ("a", "b", "c"), rows)])


class FormulaTests(unittest.TestCase):
    def test_formula_like(self):
        for s in ("=1", "+1", "-1", "@x", "\tx", "\rx", "  =x", "\x00=x"):
            self.assertTrue(is_formula_like(s), s)
        for s in ("hello", "", "a=b", " plain"):
            self.assertFalse(is_formula_like(s), s)

    def test_policy(self):
        self.assertEqual(validate_formula("=SUM(A1:A3)"), "=SUM(A1:A3)")
        for bad in ("SUM(1)", "=", '=HYPERLINK("x")', "=cmd|'/c calc'!A1", "=Sheet2!A1",
                    "=" + "1+" * 200 + "1", '=IF(A1,"x)'):
            with self.assertRaises(FormulaPolicyError, msg=bad):
                validate_formula(bad)


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.out = Path(self.dir.name) / "o.xlsx"

    def tearDown(self):
        self.dir.cleanup()

    def test_typed_cells_and_injection(self):
        res = ExcelExporter(TPL).export(wb_content(), self.out)
        ws = load_workbook(self.out)["Data"]
        self.assertEqual((ws["A2"].value, ws["A2"].data_type), ("=1+1", "s"))
        self.assertEqual((ws["B2"].value, ws["B2"].data_type), (-5, "n"))
        self.assertEqual(ws["C2"].data_type, "f")
        self.assertEqual((ws["A3"].value, ws["A3"].data_type), ("  =cmd|x", "s"))
        self.assertEqual(ws["B3"].data_type, "s")           # text "-5" stays text
        self.assertEqual(ws["C3"].data_type, "s")           # disallowed formula -> inert text
        self.assertEqual(ws["A4"].data_type, "s")
        self.assertEqual(ws["B4"].data_type, "s")
        self.assertEqual(ws["C4"].value.year, 2026)
        self.assertIs(ws["A5"].value, True)
        self.assertEqual(ws["C5"].number_format, "0%")
        self.assertTrue(any("C3" in w for w in res.warnings))
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertIn("Table1", ws.tables)

    def test_rtl_and_wide_rows_and_dup_names(self):
        wb = WorkbookContentModel(Metadata("W", "ar"), [
            Sheet("S", ("a", "a"), ((Cell("text", "1"), Cell("text", "2"), Cell("text", "3")),)),
            Sheet("s", ("x",), ())])
        res = ExcelExporter(TPL).export(wb, self.out)
        book = load_workbook(self.out)
        self.assertEqual(book.sheetnames, ["S", "s (2)"])
        self.assertTrue(book["S"].sheet_view.rightToLeft)
        self.assertEqual([c.value for c in book["S"][1]], ["a", "a (2)", "Column 3"])
        self.assertEqual(len(book["s (2)"].tables), 0)
        self.assertTrue(res.warnings)

    def test_invalid_number_becomes_text(self):
        wb = WorkbookContentModel(Metadata("W"), [Sheet("S", ("a",), ((Cell("number", "12"),),))])
        res = ExcelExporter(TPL).export(wb, self.out)
        self.assertEqual(load_workbook(self.out)["S"]["A2"].data_type, "s")
        self.assertTrue(res.warnings)


class WordTests(unittest.TestCase):
    def test_export(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.docx"
            res = WordExporter(TPL).export(doc_content("ar"), out)
            doc = Document(str(out))
            texts = [p.text for p in doc.paragraphs]
            self.assertIn("1 Intro", texts)
            self.assertIn("1.1 Sub", texts)
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(len(doc.tables[0].rows), 3)
            arabic = next(p for p in doc.paragraphs if "مرحبا" in p.text)
            self.assertIsNotNone(arabic._p.pPr.find(qn("w:bidi")))
            english = next(p for p in doc.paragraphs if p.text == "Hello world")
            self.assertIsNone(english._p.pPr.find(qn("w:bidi")) if english._p.pPr is not None else None)
            self.assertIsNotNone(doc.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")))
            self.assertTrue(res.warnings)  # wide row


class PdfTests(unittest.TestCase):
    def test_latin(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.pdf"
            content = DocumentContentModel(Metadata("Title", "en"), [
                Heading("Head", 1), Paragraph("Hello world"), ListBlock(("a",), True),
                Table(("h",), (("v",),))])
            PdfExporter(TPL).export(content, out)
            text = "".join(p.extract_text() for p in PdfReader(str(out)).pages)
            self.assertIn("Hello world", text)
            self.assertIn("Page 1", text)

    @unittest.skipUnless(FONT_DIR, "no TTF font with Arabic glyphs on this machine")
    def test_arabic_layout_with_fake_shaper(self):
        long_ar = "مرحبا بالعالم " * 80
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.pdf"
            PdfExporter(TPL, FONT_DIR, FakeShaper()).export(
                doc_content("ar").__class__(Metadata("عنوان", "ar"), [
                    Heading("عنوان", 1), Paragraph(long_ar), ListBlock(("بند", "item"), True),
                    Table(("عمود", "b"), (("١", "2"),), "جدول")]), out)
            self.assertGreaterEqual(len(PdfReader(str(out)).pages), 1)

    def test_missing_font_is_explicit(self):
        with tempfile.TemporaryDirectory() as d:
            c = DocumentContentModel(Metadata("t"), [Paragraph("مرحبا")])
            with self.assertRaises(FontNotFoundError):
                PdfExporter(TPL, Path(d), FakeShaper()).export(c, Path(d) / "o.pdf")

    def test_real_shaper_if_installed(self):
        try:
            import arabic_reshaper  # noqa: F401
            import bidi  # noqa: F401
        except ImportError:
            self.skipTest("arabic-reshaper / python-bidi not installed")
        if not FONT_DIR:
            self.skipTest("no font")
        with tempfile.TemporaryDirectory() as d:
            c = DocumentContentModel(Metadata("t", "ar"), [Paragraph("مرحبا بالعالم Mixed 123")])
            PdfExporter(TPL, FONT_DIR).export(c, Path(d) / "o.pdf")
            self.assertTrue((Path(d) / "o.pdf").stat().st_size > 0)

    def test_bad_font_in_template(self):
        bad = Template("t", "T", "d", ContentKind.DOCUMENT, "s", {}, {}, {"font": "Comic"})
        with self.assertRaises(ExportError):
            PdfExporter(bad)


class BaseBehaviourTests(unittest.TestCase):
    def test_wrong_type_and_extension(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(UnsupportedContentError):
                WordExporter(TPL).export(wb_content(), Path(d) / "o.docx")
            with self.assertRaises(ExportError):
                WordExporter(TPL).export(doc_content(), Path(d) / "o.txt")
            with self.assertRaises(ExportError):
                WordExporter(TPL).export(doc_content(), Path(d) / "nope" / "o.docx")

    def test_failed_write_keeps_existing_file_and_cleans_temp(self):
        class Boom(WordExporter):
            def _write(self, content, path, warnings):
                path.write_bytes(b"partial")
                raise RuntimeError("boom")
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.docx"
            out.write_bytes(b"old")
            with self.assertRaises(ExportError):
                Boom(TPL).export(doc_content(), out)
            self.assertEqual(out.read_bytes(), b"old")
            self.assertEqual([p.name for p in Path(d).iterdir()], ["o.docx"])

    def test_supports(self):
        self.assertTrue(ExcelExporter(TPL).supports("formulas"))
        self.assertFalse(WordExporter(TPL).supports("formulas"))
        self.assertIsInstance(WordExporter(TPL), FileExporter)


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.svc = DocumentService({OutputFormat.DOCX: WordExporter, OutputFormat.XLSX: ExcelExporter,
                                    OutputFormat.PDF: PdfExporter})

    def test_overwrite_policy(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.docx"
            self.svc.export(doc_content(), TPL, "docx", out)
            with self.assertRaises(OutputExistsError) as ctx:
                self.svc.export(doc_content(), TPL, "docx", out)
            self.assertIn("o (2).docx", ctx.exception.hint)
            self.svc.export(doc_content(), TPL, OutputFormat.DOCX, out, overwrite=True)

    def test_conversions(self):
        with tempfile.TemporaryDirectory() as d:
            r1 = self.svc.export(doc_content(), TPL, "xlsx", Path(d) / "a.xlsx")
            self.assertIn("converted", r1.warnings[0])
            book = load_workbook(Path(d) / "a.xlsx")
            self.assertEqual(book.sheetnames, ["Content", "Cap"])
            r2 = self.svc.export(wb_content(), TPL, "docx", Path(d) / "b.docx")
            self.assertIn("converted", r2.warnings[0])
            r3 = self.svc.export(wb_content(), TPL, "pdf", Path(d) / "c.pdf")
            self.assertTrue((Path(d) / "c.pdf").stat().st_size > 0)

    def test_filenames(self):
        self.assertEqual(safe_filename('a/b:c?', OutputFormat.PDF), "a_b_c_.pdf")
        self.assertEqual(safe_filename("CON", OutputFormat.PDF), "_CON.pdf")
        self.assertEqual(safe_filename("", OutputFormat.DOCX), "document.docx")
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "x.pdf"
            p.touch()
            self.assertEqual(unique_path(p).name, "x (2).pdf")


if __name__ == "__main__":
    unittest.main()
