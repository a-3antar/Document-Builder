"""Stage 6: editor logic, error paths (Ollama down, bad key, timeout, bad JSON), OCR->export flow."""
import socket
import tempfile
import time
import unittest
import urllib.error
from pathlib import Path
from unittest import mock

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from ai.gemini_provider import GeminiProvider
from ai.model_registry import ModelRegistry, ProviderRegistry
from ai.ollama_provider import OllamaProvider, urllib_json
from config.constants import OutputFormat
from core.content_editor import (EditorViewModel, block_summary, count_markers, decode_table,
                                 encode_table, kind_of)
from core.exceptions import (ModelNotAvailableError, ProviderAuthError, ProviderTimeoutError,
                             ProviderUnavailableError, TesseractNotFoundError)
from core.interfaces import OcrEngine
from core.models import (DocumentContentModel, Heading, ListBlock, Metadata, OCRResult,
                         Paragraph, Table)
from core.schema_validator import SchemaValidator
from document.excel_exporter import ExcelExporter
from document.ocr_converter import ocr_text_to_document
from document.pdf_exporter_reportlab import PdfExporter
from document.word_exporter import WordExporter
from image_processing.image_preprocessor import ImagePreprocessor
from image_processing.processing_pipeline import ProcessingPipeline
from image_processing.tesseract_ocr import TesseractEngine
from services.ai_service import AiService
from services.document_service import DocumentService
from services.review_service import ReviewService
from services.task_service import EventKind, TaskService
from tests.test_ai import BAD, GOOD, TEMPLATE, _Err, _FakeClient, _Scripted

try:
    import jsonschema  # noqa: F401
    HAS_JSONSCHEMA = True
except ImportError:
    HAS_JSONSCHEMA = False


def _final_event(tasks, func, timeout=3.0):
    """Submit work and return its terminal event (what the UI handler would receive)."""
    tid = tasks.submit(func)
    end = time.time() + timeout
    while time.time() < end:
        for ev in tasks.poll():
            if ev.task_id == tid and ev.kind is not EventKind.PROGRESS:
                return ev
        time.sleep(0.01)
    raise AssertionError("task did not finish in time")


def _ai(provider) -> AiService:
    reg = ProviderRegistry()
    reg.register("fake", lambda: provider)
    return AiService(reg, SchemaValidator())


def _down(method, url, body, timeout):
    raise ProviderUnavailableError("Ollama is not reachable.", hint="Start Ollama, then refresh.")


class _Seq(_Scripted):
    """Scripted provider whose responses may be exceptions to raise."""

    def generate_text(self, prompt, model, **options):
        item = super().generate_text(prompt, model, **options)
        if isinstance(item, Exception):
            raise item
        return item


class TableCodecTests(unittest.TestCase):
    def test_roundtrip_with_pipe(self):
        headers, rows = ("a", "b|c"), (("1", "x"), ("2", "y|z"))
        self.assertEqual(decode_table(encode_table(headers, rows)), (headers, rows))

    def test_empty_text(self):
        self.assertEqual(decode_table("  \n\n"), ((), ()))


class EditorViewModelTests(unittest.TestCase):
    def test_add_move_remove(self):
        vm = EditorViewModel(DocumentContentModel())
        self.assertEqual((vm.add("heading"), vm.add("paragraph", after=0)), (0, 1))
        self.assertEqual(vm.add("list", after=0), 1)
        self.assertEqual([kind_of(b) for b in vm.blocks], ["heading", "list", "paragraph"])
        self.assertEqual(vm.move(1, 1), 2)
        self.assertEqual(vm.move(2, 1), 2)       # edge: unchanged
        vm.remove(0)
        self.assertEqual(len(vm.blocks), 2)
        with self.assertRaises(ValueError):
            vm.add("chart")

    def test_convert_carries_text(self):
        vm = EditorViewModel(DocumentContentModel(blocks=[Paragraph("- a\n2) b")]))
        self.assertEqual(vm.convert(0, "list").items, ("a", "b"))
        self.assertEqual(vm.convert(0, "paragraph").text, "a\nb")
        self.assertEqual(vm.convert(0, "heading").text, "a b")

    def test_blocking_and_warning_issues(self):
        vm = EditorViewModel(DocumentContentModel(blocks=[
            Heading(""), ListBlock(()), Table((), ()),
            Table(("a", "b"), (("1",),)), Paragraph("x [?y?]")]))
        blocking = [i for i in vm.issues() if i.blocking]
        self.assertEqual(sorted(i.index for i in blocking), [0, 1, 2])
        warnings = [str(i) for i in vm.issues() if not i.blocking]
        self.assertTrue(any("different number" in w for w in warnings))
        self.assertTrue(any("uncertain OCR marker" in w for w in warnings))
        self.assertTrue(EditorViewModel(DocumentContentModel()).issues()[0].blocking)

    def test_summary(self):
        self.assertTrue(block_summary(Heading("T", 3)).startswith("H3"))
        self.assertIn("2 col", block_summary(Table(("a", "b"), ())))

    @unittest.skipUnless(HAS_JSONSCHEMA, "jsonschema not installed")
    def test_edited_model_is_schema_valid(self):
        vm = EditorViewModel(DocumentContentModel(Metadata("T", "mixed"), [
            Heading("H", 2), Paragraph("نص", "rtl"), ListBlock(("a",), True),
            Table(("h",), (("1",),), "cap")], needs_review=True))
        model = vm.to_model()
        self.assertFalse(model.needs_review)
        self.assertEqual(ReviewService(SchemaValidator()).validate_document(model), [])


class ProviderFailureTests(unittest.TestCase):
    def setUp(self):
        self.tasks = TaskService()

    def tearDown(self):
        self.tasks.shutdown()

    def test_ollama_stopped_reaches_ui_and_app_stays_usable(self):
        ai = _ai(OllamaProvider(http=_down))
        ev = _final_event(self.tasks, lambda ctx: ai.generate_structured("fake", "m", "x", TEMPLATE, ctx))
        self.assertIs(ev.kind, EventKind.ERROR)
        self.assertIsInstance(ev.payload, ProviderUnavailableError)
        self.assertTrue(ev.payload.hint)
        self.assertEqual(_final_event(self.tasks, lambda ctx: 7).payload, 7)

    def test_ollama_stopped_model_check_and_connection(self):
        ok, msg = OllamaProvider(http=_down).check_connection()
        self.assertFalse(ok)
        self.assertIn("Start Ollama", msg)
        reg = ProviderRegistry()
        reg.register("ollama", lambda: OllamaProvider(http=_down))
        with self.assertRaises(ProviderUnavailableError):
            ModelRegistry(reg).require("ollama", "m")

    def test_wrong_gemini_key(self):
        provider = GeminiProvider(lambda: "bad", client_factory=lambda k, t: _FakeClient(
            error=_Err(403, "denied")))
        with self.assertRaises(ProviderAuthError) as ctx:
            provider.generate_text("x", "m")
        self.assertIn("aistudio.google.com", ctx.exception.hint)
        self.assertFalse(provider.check_connection()[0])
        ai = _ai(provider)
        ev = _final_event(self.tasks, lambda ctx: ai.generate_structured("fake", "m", "x", TEMPLATE, ctx))
        self.assertIsInstance(ev.payload, ProviderAuthError)

    def test_ollama_timeouts(self):
        for err in (socket.timeout("t"), urllib.error.URLError(socket.timeout("t"))):
            with mock.patch("urllib.request.urlopen", side_effect=err):
                with self.assertRaises(ProviderTimeoutError):
                    urllib_json("GET", "http://localhost:11434/api/tags", None, 1)

    def test_ollama_model_not_installed(self):
        err = urllib.error.HTTPError("http://x", 404, "nf", {}, None)
        with mock.patch("urllib.request.urlopen", side_effect=err):
            with self.assertRaises(ModelNotAvailableError):
                urllib_json("POST", "http://localhost:11434/api/generate", {}, 1)

    def test_gemini_timeout(self):
        class ReadTimeout(Exception):
            pass

        provider = GeminiProvider(lambda: "k", client_factory=lambda k, t: _FakeClient(
            error=ReadTimeout("slow")))
        with self.assertRaises(ProviderTimeoutError):
            provider.generate_text("x", "m")


class MalformedAiOutputTests(unittest.TestCase):
    def test_truncated_json_twice_preserves_raw(self):
        prov = _Seq(['{"schema_version":"1.0","metadata":{', "still not json"])
        result = _ai(prov).generate_structured("fake", "m", "x", TEMPLATE)
        self.assertFalse(result.ok)
        self.assertIsNone(result.content)
        self.assertTrue(result.raw.startswith('{"schema_version"'))
        self.assertIn("not valid JSON", str(result.issues[0]))
        self.assertEqual(len(prov.prompts), 2)       # exactly one repair attempt

    def test_repair_call_failure_keeps_raw(self):
        prov = _Seq(["{broken", ProviderTimeoutError("slow", hint="raise the timeout")])
        result = _ai(prov).generate_structured("fake", "m", "x", TEMPLATE)
        self.assertFalse(result.ok)
        self.assertEqual(result.raw, "{broken")
        self.assertIn("Repair attempt failed", str(result.issues[-1]))

    @unittest.skipUnless(HAS_JSONSCHEMA, "jsonschema not installed")
    def test_prose_then_fenced_json_is_repaired(self):
        prov = _Seq(["Sure! Here is your document.", f"```json\n{GOOD}\n```"])
        result = _ai(prov).generate_structured("fake", "m", "x", TEMPLATE)
        self.assertTrue(result.ok)
        self.assertEqual(result.raw, "Sure! Here is your document.")

    @unittest.skipUnless(HAS_JSONSCHEMA, "jsonschema not installed")
    def test_schema_violation_twice(self):
        prov = _Seq([BAD, BAD])
        result = _ai(prov).generate_structured("fake", "m", "x", TEMPLATE)
        self.assertFalse(result.ok)
        self.assertEqual(result.raw, BAD)
        self.assertTrue(result.issues)


class _TextEngine(OcrEngine):
    id = "fake"

    def __init__(self, text):
        self._text = text

    def is_available(self):
        return True, "ok"

    def extract(self, image, languages):
        return OCRResult("", self.id, self._text)


SCAN_TEXT = ("Report title\n\n- first item\n- second item\n\nName | Qty\nApples | 3\n\n"
             "مرحبا بالعالم [?كلمة?]")


class OcrToExportFlowTests(unittest.TestCase):
    def test_ocr_review_export_all_formats(self):
        with tempfile.TemporaryDirectory() as d:
            png = Path(d) / "scan.png"
            Image.new("RGB", (20, 20), "white").save(png)
            result = ProcessingPipeline(ImagePreprocessor(), None).run(
                [png], _TextEngine(SCAN_TEXT), "ara+eng")

            vm = EditorViewModel(ocr_text_to_document(result.original_text, "Scan"))
            self.assertEqual(len(vm.blocks), 4)
            vm.convert(0, "heading")
            vm.convert(1, "list")
            vm.convert(2, "table")
            self.assertEqual([kind_of(b) for b in vm.blocks], ["heading", "list", "table", "paragraph"])
            self.assertEqual(vm.blocks[1].items, ("first item", "second item"))
            self.assertEqual((vm.blocks[2].headers, vm.blocks[2].rows),
                             (("Name", "Qty"), (("Apples", "3"),)))
            self.assertFalse([i for i in vm.issues() if i.blocking])
            self.assertEqual(count_markers(vm.blocks), 1)
            content = vm.to_model()
            if HAS_JSONSCHEMA:
                self.assertEqual(ReviewService(SchemaValidator()).validate_document(content), [])

            service = DocumentService({OutputFormat.DOCX: WordExporter,
                                       OutputFormat.XLSX: ExcelExporter,
                                       OutputFormat.PDF: PdfExporter})
            service.export(content, TEMPLATE, "docx", Path(d) / "o.docx")
            texts = [p.text for p in Document(str(Path(d) / "o.docx")).paragraphs]
            self.assertIn("Report title", texts)
            service.export(content, TEMPLATE, "xlsx", Path(d) / "o.xlsx")
            self.assertIn("Content", load_workbook(Path(d) / "o.xlsx").sheetnames)
            latin = DocumentContentModel(content.metadata, content.blocks[:3])  # no Arabic: no font needed
            service.export(latin, TEMPLATE, "pdf", Path(d) / "o.pdf")
            self.assertIn("Report title", PdfReader(str(Path(d) / "o.pdf")).pages[0].extract_text())

    def test_missing_tesseract_reaches_ui(self):
        tasks = TaskService()
        try:
            with tempfile.TemporaryDirectory() as d:
                png = Path(d) / "a.png"
                Image.new("RGB", (20, 20)).save(png)
                engine = TesseractEngine(lambda: str(Path(d) / "missing.exe"))
                pipeline = ProcessingPipeline(ImagePreprocessor(), None)
                ev = _final_event(tasks, lambda ctx: pipeline.run([png], engine, "eng"))
                self.assertIsInstance(ev.payload, TesseractNotFoundError)
                self.assertTrue(ev.payload.hint)
        finally:
            tasks.shutdown()


if __name__ == "__main__":
    unittest.main()