import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from config.constants import ContentKind, OutputFormat
from core.chart_rules import check_charts
from core.models import Cell, ChartSpec, Metadata, Sheet, WorkbookContentModel
from document.content_converter import adapt_content
from document.excel_exporter import ExcelExporter
from document.template_registry import Template

TPL = Template("t", "T", "d", ContentKind.WORKBOOK, "workbook_content.schema.json", {},
               {"header_fill": "1F4E78", "header_font_color": "FFFFFF", "default_col_width": 18,
                "freeze_panes": True}, {})


def _wb(*charts: ChartSpec, bad_numbers: bool = False) -> WorkbookContentModel:
    value = Cell("text", "n/a") if bad_numbers else Cell("number", 2)
    rows = ((Cell("text", "2021"), Cell("number", 1), Cell("number", 5)),
            (Cell("text", "2022"), value, Cell("blank")))
    return WorkbookContentModel(Metadata("W", "ar"),
                                [Sheet("Data", ("Year", "Revenue", "Cost"), rows)], charts=list(charts))


def _spec(kind="column", **over) -> ChartSpec:
    base = dict(type=kind, title="C", sheet="Data", category_column="Year",
                value_columns=("Revenue", "Cost"))
    return ChartSpec(**{**base, **over})


class ChartRulesTests(unittest.TestCase):
    def test_valid_chart(self):
        self.assertEqual(check_charts(_wb(_spec())), [])

    def test_reference_problems(self):
        for spec in (_spec(sheet="Nope"), _spec(category_column="x"),
                     _spec(value_columns=("Missing",)), _spec("pie")):
            self.assertTrue(check_charts(_wb(spec)), spec)

    def test_text_values_rejected(self):
        self.assertTrue(check_charts(_wb(_spec(), bad_numbers=True)))


class ExcelChartTests(unittest.TestCase):
    def test_native_charts_written(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.xlsx"
            res = ExcelExporter(TPL).export(
                _wb(_spec(), _spec("line"), _spec("pie", value_columns=("Revenue",))), out)
            self.assertEqual(len(load_workbook(out)["Data"]._charts), 3)
            self.assertEqual(res.warnings, ())

    def test_invalid_chart_skipped_with_warning(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.xlsx"
            res = ExcelExporter(TPL).export(_wb(_spec(sheet="Nope")), out)
            self.assertEqual(len(load_workbook(out)["Data"]._charts), 0)
            self.assertTrue(any("skipped" in w for w in res.warnings))


class ConverterTests(unittest.TestCase):
    def test_charts_note_for_non_excel(self):
        _, notes = adapt_content(_wb(_spec()), OutputFormat.PDF)
        self.assertTrue(any("Charts" in n for n in notes))


if __name__ == "__main__":
    unittest.main()